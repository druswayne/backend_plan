"""Flask-приложение ClassHub: веб-интерфейс и API для мобильного клиента."""

from __future__ import annotations

import io
import json
import os
import tempfile
from datetime import date, timedelta

from flask import Flask, Response, jsonify, redirect, render_template, request, send_file, send_from_directory, session, url_for

import auth
import database
import plan_database
import plan_routes
import plan_telegram_jobs
import service
import telegram_jobs

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
STATIC_DIR = os.path.join(BASE_DIR, "static")
TEMPLATES_DIR = os.path.join(BASE_DIR, "templates")

app = Flask(__name__, static_folder=STATIC_DIR, template_folder=TEMPLATES_DIR)
app.config["MAX_CONTENT_LENGTH"] = 64 * 1024 * 1024
app.config["SECRET_KEY"] = auth.ensure_secret_key()
app.config["PERMANENT_SESSION_LIFETIME"] = timedelta(days=14)
app.config["SESSION_COOKIE_HTTPONLY"] = True
app.config["SESSION_COOKIE_SAMESITE"] = "Lax"
app.register_blueprint(plan_routes.bp)


def json_error(exc: Exception, status: int = 400):
    message = getattr(exc, "message", None) or str(exc)
    code = getattr(exc, "status", status)
    return jsonify({"error": message}), code


@app.after_request
def add_cors(response):
    response.headers["Access-Control-Allow-Origin"] = "*"
    response.headers["Access-Control-Allow-Headers"] = "Content-Type, Authorization"
    response.headers["Access-Control-Allow-Methods"] = "GET,POST,PUT,PATCH,DELETE,OPTIONS"
    return response


@app.before_request
def protect_site():
    return auth.require_login()


@app.route("/api/<path:_any>", methods=["OPTIONS"])
def api_options(_any):
    return ("", 204)


@app.get("/")
def index():
    return send_from_directory(TEMPLATES_DIR, "index.html")


@app.route("/login", methods=["GET", "POST"])
def login():
    error = None
    nxt = request.values.get("next") or "/"
    if request.method == "POST":
        data = request.get_json(silent=True) if request.is_json else None
        username = (data or request.form).get("username") or ""
        password = (data or request.form).get("password") or ""
        nxt = (data or request.form).get("next") or nxt
        if auth.verify_password(username, password):
            auth.login_user(username)
            if request.is_json:
                return jsonify({"ok": True})
            if not nxt.startswith("/"):
                nxt = "/"
            return redirect(nxt)
        error = "Неверный логин или пароль"
        if request.is_json:
            return jsonify({"error": error}), 401
    return render_template("login.html", error=error, next=nxt, username=auth.current_username())


@app.get("/logout")
def logout():
    auth.logout_user()
    return redirect(url_for("login"))


@app.get("/api/auth/me")
def api_auth_me():
    return jsonify({"username": session.get("user") or auth.current_username()})


@app.put("/api/auth/credentials")
def api_auth_credentials():
    data = request.get_json(force=True, silent=True) or {}
    current_password = data.get("currentPassword") or data.get("current_password") or ""
    username = (data.get("username") or auth.current_username()).strip()
    new_password = data.get("password") or data.get("newPassword") or ""
    actor = session.get("user") or (request.authorization.username if request.authorization else "")
    if not auth.verify_password(actor or username, current_password):
        return jsonify({"error": "Текущий пароль неверный"}), 400
    try:
        auth.save_credentials(username, new_password)
    except ValueError as exc:
        return jsonify({"error": str(exc)}), 400
    auth.login_user(username)
    return jsonify({"ok": True, "username": username})


@app.get("/api/health")
def health():
    return jsonify({"ok": True, "app": "ClassHub"})


@app.get("/api/home")
def api_home():
    try:
        return jsonify(service.dashboard(request.args.get("month")))
    except service.AppError as exc:
        return json_error(exc)


@app.get("/api/students")
def api_students():
    archived = request.args.get("archived", "0") in ("1", "true", "yes")
    return jsonify(service.list_students(archived=archived))


@app.post("/api/students")
def api_add_student():
    data = request.get_json(force=True, silent=True) or {}
    try:
        student_id = service.add_student(
            data.get("name", ""),
            data.get("grade", ""),
            data.get("subjects") or [],
            data.get("address"),
        )
        return jsonify({"id": student_id}), 201
    except service.AppError as exc:
        return json_error(exc)


@app.get("/api/students/<int:student_id>")
def api_student(student_id: int):
    try:
        return jsonify(service.student_detail(student_id, request.args.get("month")))
    except service.AppError as exc:
        return json_error(exc)


@app.put("/api/students/<int:student_id>")
def api_update_student(student_id: int):
    data = request.get_json(force=True, silent=True) or {}
    try:
        service.update_student(student_id, data.get("name", ""), data.get("grade", ""), data.get("address"))
        return jsonify({"ok": True})
    except service.AppError as exc:
        return json_error(exc)


@app.post("/api/students/<int:student_id>/archive")
def api_archive_student(student_id: int):
    data = request.get_json(force=True, silent=True) or {}
    archived = data.get("archived", True)
    service.set_student_archived(student_id, bool(archived))
    return jsonify({"ok": True})


@app.delete("/api/students/<int:student_id>")
def api_delete_student(student_id: int):
    service.delete_student(student_id)
    return jsonify({"ok": True})


@app.post("/api/students/<int:student_id>/subjects")
def api_add_subject(student_id: int):
    data = request.get_json(force=True, silent=True) or {}
    try:
        subject_id = service.add_subject(student_id, data.get("name", ""), float(data.get("price") or 0))
        return jsonify({"id": subject_id}), 201
    except service.AppError as exc:
        return json_error(exc)


@app.put("/api/subjects/<int:subject_id>")
def api_update_subject(subject_id: int):
    data = request.get_json(force=True, silent=True) or {}
    try:
        service.update_subject(subject_id, data.get("name", ""), float(data.get("price") or 0))
        return jsonify({"ok": True})
    except service.AppError as exc:
        return json_error(exc)


@app.delete("/api/subjects/<int:subject_id>")
def api_delete_subject(subject_id: int):
    try:
        service.delete_subject(subject_id)
        return jsonify({"ok": True})
    except service.AppError as exc:
        return json_error(exc)


@app.get("/api/lessons")
def api_lessons():
    from_day = request.args.get("from")
    to_day = request.args.get("to")
    if from_day and to_day:
        start, end = int(from_day), int(to_day)
    else:
        start_date = date.fromisoformat(request.args.get("start") or date.today().isoformat())
        end_date = date.fromisoformat(request.args.get("end") or start_date.isoformat())
        start, end = service.to_epoch_day(start_date), service.to_epoch_day(end_date)
    include_archived = (
        request.args.get("includeArchived", request.args.get("includeArchived", "0"))
        in ("1", "true", "yes")
    )
    upcoming_only = request.args.get("upcoming", "0") in ("1", "true", "yes")
    return jsonify(
        service.lessons_in_range(
            start, end, upcoming_only=upcoming_only, include_archived=include_archived
        )
    )


@app.get("/api/lessons/<int:occurrence_id>")
def api_lesson(occurrence_id: int):
    try:
        return jsonify(service.get_lesson(occurrence_id))
    except service.AppError as exc:
        return json_error(exc)


@app.get("/api/series/<int:series_id>")
def api_series(series_id: int):
    try:
        return jsonify(service.get_series(series_id))
    except service.AppError as exc:
        return json_error(exc)


@app.post("/api/lessons")
def api_create_lesson():
    data = request.get_json(force=True, silent=True) or {}
    try:
        series_id = service.create_lesson(data)
        return jsonify({"seriesId": series_id}), 201
    except service.ConflictError as exc:
        return json_error(exc)
    except service.AppError as exc:
        return json_error(exc)


@app.patch("/api/lessons/<int:occurrence_id>/status")
def api_lesson_status(occurrence_id: int):
    data = request.get_json(force=True, silent=True) or {}
    try:
        service.change_status(occurrence_id, data.get("status", ""))
        return jsonify({"ok": True})
    except (service.ConflictError, service.AppError) as exc:
        return json_error(exc)


@app.patch("/api/lessons/<int:occurrence_id>/notes")
def api_lesson_notes(occurrence_id: int):
    data = request.get_json(force=True, silent=True) or {}
    service.update_notes(occurrence_id, data.get("notes") or "")
    return jsonify({"ok": True})


@app.patch("/api/lessons/<int:occurrence_id>/reminder")
def api_lesson_reminder(occurrence_id: int):
    data = request.get_json(force=True, silent=True) or {}
    service.update_reminder(occurrence_id, data.get("reminderMinutes"))
    return jsonify({"ok": True})


@app.patch("/api/lessons/<int:occurrence_id>/reschedule")
def api_lesson_reschedule(occurrence_id: int):
    data = request.get_json(force=True, silent=True) or {}
    try:
        service.reschedule_occurrence(occurrence_id, data)
        return jsonify({"ok": True})
    except (service.ConflictError, service.AppError) as exc:
        return json_error(exc)


@app.patch("/api/lessons/<int:occurrence_id>/subject")
def api_lesson_subject(occurrence_id: int):
    data = request.get_json(force=True, silent=True) or {}
    try:
        service.change_subject(
            occurrence_id,
            int(data.get("studentSubjectId") or 0),
            bool(data.get("applyToFuture")),
        )
        return jsonify({"ok": True})
    except service.AppError as exc:
        return json_error(exc)


@app.post("/api/series/<int:series_id>/future")
def api_series_future(series_id: int):
    data = request.get_json(force=True, silent=True) or {}
    try:
        service.update_series_future(series_id, data)
        return jsonify({"ok": True})
    except (service.ConflictError, service.AppError) as exc:
        return json_error(exc)


@app.post("/api/series/<int:series_id>/extend")
def api_series_extend(series_id: int):
    data = request.get_json(force=True, silent=True) or {}
    try:
        service.extend_series(series_id, data.get("endDate") or "")
        return jsonify({"ok": True})
    except service.AppError as exc:
        return json_error(exc)


@app.delete("/api/lessons/<int:occurrence_id>")
def api_delete_lesson(occurrence_id: int):
    service.delete_occurrence(occurrence_id)
    return jsonify({"ok": True})


@app.delete("/api/series/<int:series_id>")
def api_delete_series(series_id: int):
    service.delete_series(series_id)
    return jsonify({"ok": True})


@app.get("/api/journals")
def api_journals():
    return jsonify(service.journal_summaries())


@app.get("/api/journals/<int:student_id>")
def api_student_journal(student_id: int):
    try:
        return jsonify(service.student_journal(student_id))
    except service.AppError as exc:
        return json_error(exc)


@app.get("/api/payments")
def api_payments():
    return jsonify(service.list_payments())


@app.get("/api/payments/unpaid")
def api_unpaid_students():
    return jsonify(service.list_unpaid_students())


@app.post("/api/payments")
def api_add_payment():
    data = request.get_json(force=True, silent=True) or {}
    try:
        payment_id = service.process_payment(
            int(data.get("studentId") or 0),
            float(data.get("amount") or 0),
            data.get("note") or "",
        )
        return jsonify({"id": payment_id}), 201
    except service.AppError as exc:
        return json_error(exc)


@app.delete("/api/payments/<int:payment_id>")
def api_delete_payment(payment_id: int):
    service.delete_payment(payment_id)
    return jsonify({"ok": True})


@app.get("/api/settings")
def api_get_settings():
    try:
        return jsonify(service.get_settings())
    except service.AppError as exc:
        return json_error(exc)


@app.put("/api/settings")
def api_update_settings():
    data = request.get_json(force=True, silent=True) or {}
    try:
        return jsonify(service.update_settings(data))
    except service.AppError as exc:
        return json_error(exc)


@app.get("/api/backup")
def api_download_backup():
    tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".db")
    tmp.close()
    try:
        database.export_copy(tmp.name)
        with open(tmp.name, "rb") as fh:
            payload = io.BytesIO(fh.read())
        payload.seek(0)
        filename = f"classhub_backup_{date.today().isoformat()}.db"
        return send_file(
            payload,
            as_attachment=True,
            download_name=filename,
            mimetype="application/octet-stream",
        )
    finally:
        try:
            os.remove(tmp.name)
        except OSError:
            pass


@app.put("/api/backup")
@app.post("/api/backup")
def api_upload_backup():
    uploaded = request.files.get("file") or request.files.get("database")
    raw = None
    if uploaded is not None:
        raw = uploaded.read()
    elif request.data:
        raw = request.data
    if not raw:
        return jsonify({"error": "Файл базы данных не передан"}), 400
    tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".db")
    try:
        tmp.write(raw)
        tmp.close()
        database.replace_database(tmp.name)
        return jsonify({"ok": True, "bytes": len(raw)})
    except ValueError as exc:
        return json_error(exc, 400)
    finally:
        try:
            os.remove(tmp.name)
        except OSError:
            pass


@app.get("/api/export/schedule")
def api_export_schedule():
    week_iso = request.args.get("week") or date.today().isoformat()
    anchor = date.fromisoformat(week_iso)
    start, end = service.week_bounds(anchor)
    lessons = service.lessons_in_range(start, end)
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
    except ImportError:
        return jsonify({"error": "Для экспорта расписания установите openpyxl"}), 500

    monday = anchor - timedelta(days=anchor.isoweekday() - 1)
    wb = Workbook()
    sheet = wb.active
    sheet.title = "Расписание"
    headers = ["Дата", "День", "Время", "Конец", "Ученик", "Предмет", "Класс", "Место", "Статус"]
    sheet.append(headers)
    status_map = {
        "EXPECTED": "Запланировано",
        "CONDUCTED": "Проведено",
        "CANCELLED": "Отменено",
    }
    weekdays = ["Пн", "Вт", "Ср", "Чт", "Пт", "Сб", "Вс"]
    for lesson in lessons:
        occ = lesson["occurrence"]
        day = service.from_epoch_day(occ["dateEpochDay"])
        start_m = occ["startTimeMinutes"]
        end_m = start_m + occ["durationMinutes"]
        place = "Онлайн" if occ["locationType"] == "ONLINE" else (occ["address"] or "Офлайн")
        sheet.append(
            [
                day.strftime("%d.%m.%Y"),
                weekdays[day.isoweekday() - 1],
                f"{start_m // 60:02d}:{start_m % 60:02d}",
                f"{end_m // 60:02d}:{end_m % 60:02d}",
                lesson["student"]["name"],
                lesson["subjectName"],
                lesson["student"].get("grade") or "",
                place,
                status_map.get(occ["status"], occ["status"]),
            ]
        )
    header_fill = PatternFill("solid", fgColor="4F46E5")
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
    bio = io.BytesIO()
    wb.save(bio)
    bio.seek(0)
    filename = f"classhub_schedule_{monday.isoformat()}.xlsx"
    return send_file(
        bio,
        as_attachment=True,
        download_name=filename,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


@app.post("/api/telegram/daily")
def api_telegram_daily():
    data = request.get_json(force=True, silent=True) or {}
    test = bool(data.get("test") or request.args.get("test") in ("1", "true", "yes"))
    try:
        return jsonify(telegram_jobs.send_daily_schedule(test=test))
    except service.AppError as exc:
        return json_error(exc)


@app.get("/api/telegram/status")
def api_telegram_status():
    return jsonify(telegram_jobs.status())


@app.post("/api/telegram/backup")
def api_telegram_backup():
    settings = service.get_settings()
    tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".db")
    tmp.close()
    try:
        database.export_copy(tmp.name)
        with open(tmp.name, "rb") as fh:
            content = fh.read()
        filename = f"classhub_backup_{date.today().isoformat()}.db"
        caption = f"Резервная копия ClassHub\nДата: {date.today().strftime('%d.%m.%Y')}"
        token = settings.get("telegramBotToken") or settings.get("telegramBotToken") or ""
        chat_id = settings.get("telegramChatId") or settings.get("telegramChatId") or ""
        telegram_jobs.send_document(token, chat_id, filename, content, caption)
        return jsonify({"ok": True})
    except service.AppError as exc:
        return json_error(exc)
    finally:
        try:
            os.remove(tmp.name)
        except OSError:
            pass


def create_app() -> Flask:
    database.init_db()
    plan_database.init_db()
    telegram_jobs.start_scheduler()
    plan_telegram_jobs.start_scheduler()
    return app


if __name__ == "__main__":
    debug = True
    # Reloader-родитель не должен держать SQLite, иначе замена файла базы не сработает.
    if os.environ.get("WERKZEUG_RUN_MAIN") == "true" or not debug:
        database.init_db()
        plan_database.init_db()
        telegram_jobs.start_scheduler()
        plan_telegram_jobs.start_scheduler()
    app.run(host="0.0.0.0", port=int(os.environ.get("PORT", "5000")), debug=debug)
